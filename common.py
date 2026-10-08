"""
Shared helpers used by every stage of the pipeline.

Design notes
------------
* Trip.com renders reviews client-side and its review API rejects replayed/synthetic calls
  (HTTP 430). The crawler therefore never calls the API itself: it opens the public page in a
  real browser, *listens* to the JSON responses the page requests, and parses those.
* Field names in Trip.com's JSON change over time, so parsing is key-name tolerant: each output
  field has an ordered list of candidate keys (see REVIEW_* constants below).
* CAPTCHA / login walls are detected and reported, never solved or bypassed.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import json
import logging
import os
import random
import re
import time
from pathlib import Path
from typing import Any, Iterable, Iterator

import pandas as pd

import config as C

# ============================================================== setup / IO

def ensure_dirs() -> None:
    for d in (C.INPUT_DIR, C.OUTPUT_DIR, C.STAGE_DIR, C.WORK_DIR, C.DEBUG_DIR, C.META_CACHE_DIR,
              C.REVIEW_DIR, C.RAW_RESPONSE_DIR, C.IMAGE_DIR):
        d.mkdir(parents=True, exist_ok=True)


def get_logger(name: str) -> logging.Logger:
    ensure_dirs()
    logger = logging.getLogger(name)
    if logger.handlers:
        return logger
    logger.setLevel(logging.INFO)
    fmt = logging.Formatter("%(asctime)s | %(name)s | %(levelname)s | %(message)s")
    fh = logging.FileHandler(C.WORK_DIR / "pipeline.log", encoding="utf-8")
    fh.setFormatter(fmt)
    sh = logging.StreamHandler()
    sh.setFormatter(fmt)
    logger.addHandler(fh)
    logger.addHandler(sh)
    return logger


def find_input(name: str) -> Path:
    """Look for a source workbook in inputs/ first, then the project root."""
    for p in (C.INPUT_DIR / name, C.ROOT / name):
        if p.exists():
            return p
    raise FileNotFoundError(f"Input file not found: {name} (put it in {C.INPUT_DIR})")


def read_json(path: Path, default: Any = None) -> Any:
    try:
        return json.loads(Path(path).read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return default


def write_json_atomic(path: Path, obj: Any) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(obj, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    os.replace(tmp, path)


def now_iso() -> str:
    return dt.datetime.now().astimezone().isoformat(timespec="seconds")


def load_run_config() -> dict:
    """Freeze stay conditions on first run so resumed runs use identical dates."""
    cfg = read_json(C.RUN_CONFIG)
    if cfg:
        return cfg
    checkin = dt.date.today() + dt.timedelta(days=C.DEFAULT_CHECKIN_OFFSET_DAYS)
    checkout = checkin + dt.timedelta(days=C.STAY_NIGHTS)
    cfg = {
        "checkin": checkin.isoformat(),
        "checkout": checkout.isoformat(),
        "adults": C.ADULTS,
        "children": C.CHILDREN,
        "rooms": C.ROOMS,
        "frozen_at": now_iso(),
    }
    write_json_atomic(C.RUN_CONFIG, cfg)
    return cfg


def price_basis_text(cfg: dict) -> str:
    # 2026-10: the dated detail page redirects automated anonymous browsers to the sign-in page (not
    # bypassed), so the price is the one Trip.com prints on the public review page for every hotel.
    return ("Trip.com public review page 'displayPrice' (lowest nightly rate before tax, USD) for the "
            "default check-in date shown to an anonymous visitor on the crawl day (date stored as "
            "price_checkin in work/hotel_meta/<code>.json); same source and currency for every hotel")


def polite_sleep(rng: tuple[float, float]) -> None:
    time.sleep(random.uniform(*rng))

# ============================================================== hotel identity helpers

_TRIP_CODE_RE = re.compile(r"hotel-detail-(\d+)")


def trip_code_from_url(url: Any) -> str | None:
    if not isinstance(url, str):
        return None
    m = _TRIP_CODE_RE.search(url)
    return m.group(1) if m else None


def review_url_from_detail(url: str) -> str:
    url = url.split("?")[0].rstrip("/")
    return url if url.endswith("review.html") else url + "/review.html"


def detail_url_from_code(code: str, city: str = "city", slug: str = "hotel") -> str:
    return f"{C.TRIP_BASE}/hotels/{city.lower().replace(' ', '-')}-hotel-detail-{code}/{slug}/"


def brand_of(name: str) -> tuple[str, str]:
    low = (name or "").lower()
    for pat, brand, group in C.BRAND_PATTERNS:
        if re.search(pat, low):
            return brand, group
    return "Independent/Other", "Independent/Other"


def area_tokens(name: str) -> set[str]:
    low = (name or "").lower()
    return {t for t in C.AREA_TOKENS if re.search(rf"\b{t}\b", low)}


def clean_str(x: Any) -> str:
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return ""
    return str(x).strip()

# ============================================================== Excel output

def write_excel(path: Path, sheets: dict[str, pd.DataFrame], notes: list[tuple[str, str]] | None = None) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name[:31], index=False)
        if notes:
            pd.DataFrame(notes, columns=["item", "detail"]).to_excel(xw, sheet_name="Notes", index=False)
    wb = load_workbook(path)
    head_fill = PatternFill("solid", start_color="DDE4EE")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                cell.font = Font(name="Arial", size=10, bold=(cell.row == 1))
                if cell.row == 1:
                    cell.fill = head_fill
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            letter = col[0].column_letter
            width = max((len(str(c.value)) if c.value is not None else 0) for c in col[:300])
            ws.column_dimensions[letter].width = max(8, min(60, width + 2))
    wb.save(path)

# ============================================================== JSON walking

def iter_dicts(obj: Any, _depth: int = 0, max_depth: int = 40) -> Iterator[dict]:
    if _depth > max_depth:
        return
    if isinstance(obj, dict):
        yield obj
        for v in obj.values():
            yield from iter_dicts(v, _depth + 1, max_depth)
    elif isinstance(obj, list):
        for v in obj:
            yield from iter_dicts(v, _depth + 1, max_depth)


def iter_lists(obj: Any, _depth: int = 0, max_depth: int = 40) -> Iterator[tuple[str, list]]:
    """Yield (parent_key, list) for every list in the structure."""
    if _depth > max_depth:
        return
    if isinstance(obj, dict):
        for k, v in obj.items():
            if isinstance(v, list):
                yield k, v
            yield from iter_lists(v, _depth + 1, max_depth)
    elif isinstance(obj, list):
        for v in obj:
            yield from iter_lists(v, _depth + 1, max_depth)


def first_scalar(d: dict, keys: Iterable[str], depth: int = 0, skip: set[str] | None = None) -> Any:
    """First non-empty scalar found under any of `keys` (searching nested dicts up to `depth`)."""
    skip = skip or set()
    if not isinstance(d, dict):
        return None
    for k in keys:
        v = d.get(k)
        if v not in (None, "", [], {}) and not isinstance(v, (dict, list)):
            return v
    if depth > 0:
        for k, v in d.items():
            if k in skip or not isinstance(v, dict):
                continue
            r = first_scalar(v, keys, depth - 1, skip)
            if r not in (None, ""):
                return r
    return None


def to_number(x: Any) -> float | None:
    if x is None or isinstance(x, bool):
        return None
    if isinstance(x, (int, float)):
        return None if (isinstance(x, float) and x != x) else float(x)     # NaN → None
    m = re.search(r"-?\d+(?:[.,]\d+)?", str(x).replace(",", ""))
    return float(m.group(0)) if m else None


def normalize_date(v: Any) -> str:
    if v in (None, ""):
        return ""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        ts = float(v)
        if ts > 1e12:
            ts /= 1000.0
        if ts > 1e9:
            return dt.datetime.fromtimestamp(ts, dt.timezone.utc).isoformat(timespec="seconds")
        return str(v)
    s = str(v).strip()
    m = re.match(r"/Date\((\d+)([+-]\d{4})?\)/", s)
    if m:
        return dt.datetime.fromtimestamp(int(m.group(1)) / 1000, dt.timezone.utc).isoformat(timespec="seconds")
    if re.fullmatch(r"\d{12,13}", s):
        return normalize_date(int(s))
    return s

# ============================================================== review parsing

REVIEW_ID_KEYS = ["commentId", "reviewId", "id"]
REVIEW_TEXT_KEYS = ["originalContent", "originContent", "content", "commentContent", "reviewContent", "text"]
REVIEW_TRANSLATED_KEYS = ["translatedContent", "translateContent", "contentTranslated", "translation"]
REVIEW_TRANSLATED_FLAG_KEYS = ["isTranslated", "translated", "isTranslate"]
REVIEW_RATING_KEYS = ["ratingAll", "rating", "score", "commentScore", "totalScore", "overallRating", "point"]
REVIEW_DATE_KEYS = ["createDate", "commentDate", "reviewDate", "publishDate", "postDate", "createTime", "commentTime"]
REVIEW_STAY_KEYS = ["checkInDate", "checkinDate", "stayDate", "checkInTime"]
REVIEW_TRAVEL_KEYS = ["travelTypeText", "travelTypeName", "travelType", "tripType", "travelPurpose"]
REVIEW_ROOM_KEYS = ["roomName", "roomTypeName", "baseRoomName", "roomType"]
REVIEW_TITLE_KEYS = ["title", "commentTitle", "reviewTitle"]
REVIEW_LANG_KEYS = ["language", "lang", "languageCode", "commentLanguage", "locale"]
REVIEW_HELPFUL_KEYS = ["usefulCount", "helpfulCount", "usefulNum", "likeCount", "likes"]
USER_KEYS = ["userInfo", "user", "reviewer", "author", "userProfile", "memberInfo"]
USER_NAME_KEYS = ["nickName", "nickname", "userName", "displayName", "reviewerName", "name"]
USER_COUNTRY_KEYS = ["countryName", "regionName", "country", "region", "userRegion", "nationality", "location"]
IMAGE_LIST_KEYS = ["imageList", "images", "pictureList", "photoList", "imgList", "commentImages",
                   "pictures", "imageInfos", "photos", "mediaList", "media"]
IMAGE_URL_KEYS = ["originalUrl", "originUrl", "bigUrl", "largeUrl", "fullUrl", "url", "imageUrl",
                  "picUrl", "src", "middleUrl", "smallUrl", "thumbUrl"]
REPLY_KEYS = ["feedbackList", "replyList", "hotelReply", "reply", "replies", "responseList"]
NON_REVIEW_SUBTREES = set(IMAGE_LIST_KEYS) | set(REPLY_KEYS)


def _is_review_like(d: dict) -> bool:
    if not isinstance(d, dict):
        return False
    has_id = any(isinstance(d.get(k), (str, int)) and d.get(k) not in ("", 0) for k in REVIEW_ID_KEYS)
    text = next((d.get(k) for k in REVIEW_TEXT_KEYS if isinstance(d.get(k), str)), None)
    has_ctx = (first_scalar(d, REVIEW_RATING_KEYS, 2, NON_REVIEW_SUBTREES) is not None
               or any(k in d for k in REVIEW_DATE_KEYS) or any(k in d for k in USER_KEYS))
    return bool(has_id and text is not None and has_ctx)


def _image_urls(d: dict) -> list[str]:
    urls: list[str] = []
    containers = [d] + [v for k, v in d.items() if isinstance(v, dict) and k not in REPLY_KEYS]
    for cont in containers:
        for key in IMAGE_LIST_KEYS:
            lst = cont.get(key)
            if not isinstance(lst, list):
                continue
            for item in lst:
                u = None
                if isinstance(item, str):
                    u = item
                elif isinstance(item, dict):
                    u = first_scalar(item, IMAGE_URL_KEYS, 1)
                    if item.get("type") in ("video", 2) and not u:
                        continue
                if isinstance(u, str) and u.strip():
                    u = u.strip()
                    if u.startswith("//"):
                        u = "https:" + u
                    elif u.startswith("/"):                 # CDN-relative review photo path
                        u = C.REVIEW_IMAGE_BASE + u
                    if u not in urls:
                        urls.append(u)
        if urls:
            break
    return urls


def _reply_text(d: dict) -> str:
    for k in REPLY_KEYS:
        v = d.get(k)
        if isinstance(v, list) and v:
            texts = [first_scalar(x, ["content", "replyContent", "text"]) for x in v if isinstance(x, dict)]
            texts = [t for t in texts if isinstance(t, str)]
            if texts:
                return "\n---\n".join(texts)
        if isinstance(v, dict):
            t = first_scalar(v, ["content", "replyContent", "text"])
            if isinstance(t, str):
                return t
        if isinstance(v, str) and v.strip():
            return v
    return ""


def _pseudonym(name: str) -> str:
    if not name:
        return ""
    return "U_" + hashlib.sha256((C.PSEUDONYM_SALT + name).encode("utf-8")).hexdigest()[:12]


def normalize_review(d: dict) -> dict:
    """Map one Trip.com review object to the study schema. Text is never altered."""
    rid = first_scalar(d, REVIEW_ID_KEYS)
    original = next((d.get(k) for k in REVIEW_TEXT_KEYS if isinstance(d.get(k), str)), "")
    has_original_key = any(isinstance(d.get(k), str) for k in ("originalContent", "originContent"))
    translated_flag = any(bool(d.get(k)) for k in REVIEW_TRANSLATED_FLAG_KEYS)
    if has_original_key:
        text_source = "original"
    elif translated_flag:
        text_source = "possibly_translated"      # page flagged translation and no original field
    else:
        text_source = "original"
    user = next((d.get(k) for k in USER_KEYS if isinstance(d.get(k), dict)), {}) or {}
    name = clean_str(first_scalar(user, USER_NAME_KEYS) or first_scalar(d, ["nickName", "userName", "reviewerName"]))
    if C.PSEUDONYMIZE_REVIEWER_NAME:
        name = _pseudonym(name)
    country = clean_str(first_scalar(user, USER_COUNTRY_KEYS, 1) or first_scalar(d, USER_COUNTRY_KEYS))
    extra = {k: v for k, v in user.items()
             if not isinstance(v, (dict, list)) and k not in USER_NAME_KEYS + USER_COUNTRY_KEYS + ["avatar", "headPhoto", "uid", "userId"]}
    rating = to_number(first_scalar(d, REVIEW_RATING_KEYS, 2, NON_REVIEW_SUBTREES))
    imgs = _image_urls(d)
    return {
        "review_id": clean_str(rid),
        "review_text": original if isinstance(original, str) else "",
        "text_source": text_source,
        "review_title": clean_str(first_scalar(d, REVIEW_TITLE_KEYS)),
        "review_rating": rating,
        "rating_scale": (to_number(d.get("ratingFull")) or (10 if (rating is not None and rating > 5) else 5))
                        if rating is not None else None,
        "review_date": normalize_date(first_scalar(d, REVIEW_DATE_KEYS)),
        "stay_date": normalize_date(first_scalar(d, REVIEW_STAY_KEYS)),
        "traveler_type": clean_str(first_scalar(d, REVIEW_TRAVEL_KEYS, 1, NON_REVIEW_SUBTREES)),
        "room_type": clean_str(first_scalar(d, REVIEW_ROOM_KEYS, 1, NON_REVIEW_SUBTREES)),
        "reviewer_name": name,
        "reviewer_country": country,
        "language": clean_str(first_scalar(d, REVIEW_LANG_KEYS)),
        "helpful_count": to_number(first_scalar(d, REVIEW_HELPFUL_KEYS)),
        "hotel_reply_text": _reply_text(d),
        "reviewer_extra": json.dumps(extra, ensure_ascii=False) if extra else "",
        "image_urls_all": imgs,
    }


def parse_reviews(obj: Any) -> list[dict]:
    """Find review objects anywhere in a JSON payload and normalise them."""
    found: dict[str, dict] = {}
    for _key, lst in iter_lists(obj):
        if _key in NON_REVIEW_SUBTREES:          # hotel replies / image lists are not reviews
            continue
        dicts = [x for x in lst if isinstance(x, dict)]
        if not dicts:
            continue
        hits = [x for x in dicts if _is_review_like(x)]
        if len(hits) >= max(1, len(dicts) // 2):
            for h in hits:
                r = normalize_review(h)
                if r["review_id"] and r["review_id"] not in found:
                    found[r["review_id"]] = r
    return list(found.values())


def parse_total_review_count(obj: Any) -> int | None:
    keys = ["totalCount", "commentCount", "totalReviews", "reviewCount", "commentTotal", "total"]
    list_keys = ("groupList", "commentList", "reviewList", "comments", "reviews")
    for d in iter_dicts(obj):            # 1) the dict that actually carries the review list
        if any(isinstance(d.get(x), list) for x in list_keys):
            for k in keys:
                v = d.get(k)
                if isinstance(v, (int, float)) and not isinstance(v, bool) and v >= 0:
                    return int(v)
    for d in iter_dicts(obj):            # 2) rating header (commentRating.showCommentNum)
        v = d.get("showCommentNum")
        if isinstance(v, (int, float)) and not isinstance(v, bool) and v >= 0:
            return int(v)
    for d in iter_dicts(obj):
        for k in keys:
            v = d.get(k)
            if isinstance(v, (int, float)) and not isinstance(v, bool) and v >= 0:
                if any(isinstance(d.get(x), list) for x in ("commentList", "reviewList", "comments", "reviews")) or k != "total":
                    return int(v)
    return None

# ============================================================== hotel meta parsing

HOTEL_ID_KEYS = ["hotelId", "masterHotelId", "hotelID", "id"]
STAR_KEYS = ["star", "starLevel", "hotelStar", "starRating", "stars"]
HOTEL_RATING_KEYS = ["score", "commentScore", "hotelScore", "overallScore", "rating", "ratingAll"]
HOTEL_COUNT_KEYS = ["commentCount", "reviewCount", "totalReviews", "commentTotal", "reviewsCount"]
PRICE_KEYS = ["displayPrice", "price", "minPrice", "salePrice", "avgPrice", "amount", "priceValue"]
CURRENCY_KEYS = ["currency", "currencyCode", "curr"]


def _hotel_subtrees(blobs: list, code: str | None) -> list[dict]:
    if not code:
        return []
    out = []
    for b in blobs:
        for d in iter_dicts(b):
            if any(str(d.get(k)) == str(code) for k in HOTEL_ID_KEYS):
                out.append(d)
    return out


def _search_keys(dicts: list[dict], keys: list[str], valid) -> Any:
    for root in dicts:
        for d in iter_dicts(root, max_depth=4):
            for k in keys:
                v = d.get(k)
                if v is not None and not isinstance(v, (dict, list, bool)) and valid(v):
                    return v
    return None


def parse_hotel_meta(blobs: list, page_text: str, code: str | None) -> dict:
    subs = _hotel_subtrees(blobs, code)
    source = "hotel_subtree" if subs else "global"
    roots = subs or blobs
    star = _search_keys(roots, STAR_KEYS, lambda v: (to_number(v) or 0) > 0 and to_number(v) <= 5)
    rating = _search_keys(roots, HOTEL_RATING_KEYS, lambda v: (to_number(v) or 0) > 0 and to_number(v) <= 10)
    count = _search_keys(roots, HOTEL_COUNT_KEYS, lambda v: (to_number(v) or -1) >= 0)
    city_id = _search_keys(roots, ["cityId", "cityID"], lambda v: (to_number(v) or 0) > 0)
    city_name = _search_keys(roots, ["cityName", "cityEnName"], lambda v: isinstance(v, str))
    name = _search_keys(subs, ["hotelName", "hotelEnName", "name"], lambda v: isinstance(v, str)) if subs else None
    # DOM fallbacks
    if count is None:
        m = re.search(r"([\d,]+)\s+(?:reviews|Reviews|条点评|件の口コミ|개의? 리뷰)", page_text or "")
        if m:
            count = int(m.group(1).replace(",", ""))
    if rating is None:
        m = re.search(r"\b(\d(?:\.\d)?)\s*/\s*(10|5)\b", page_text or "")
        if m:
            rating = float(m.group(1))
    # Price: minimum displayed price among dicts that carry both a price and a currency
    prices = []
    for b in blobs:
        for d in iter_dicts(b):
            cur = first_scalar(d, CURRENCY_KEYS)
            if not isinstance(cur, str) or not re.fullmatch(r"[A-Z]{3}", cur):
                continue
            for k in PRICE_KEYS:
                p = to_number(d.get(k)) if not isinstance(d.get(k), (dict, list)) else None
                if p and p > 0:
                    prices.append((p, cur))
    price_value, price_currency = (None, None)
    if prices:
        curr_counts = pd.Series([c for _, c in prices]).value_counts()
        main_cur = curr_counts.index[0]
        price_value = min(p for p, c in prices if c == main_cur)
        price_currency = main_cur
    rating_n = to_number(rating)
    return {
        "star_rating": to_number(star),
        "overall_rating": rating_n,
        "overall_rating_scale": (10 if rating_n and rating_n > 5 else (5 if rating_n else None)),
        "total_review_count": int(to_number(count)) if count is not None else None,
        "price_value": price_value,
        "price_currency": price_currency,
        "trip_city_id": int(to_number(city_id)) if city_id else None,
        "trip_city_name": city_name,
        "trip_hotel_name": name,
        "meta_source": source,
    }


def extract_facility_text(blobs: list, dom_facility_text: str) -> str:
    parts = [dom_facility_text or ""]
    for b in blobs:
        for d in iter_dicts(b):
            for k, v in d.items():
                if re.search(r"facilit|amenit|feature|policy|service", k, re.I):
                    if isinstance(v, str):
                        parts.append(v)
                    elif isinstance(v, list):
                        for item in v:
                            if isinstance(item, str):
                                parts.append(item)
                            elif isinstance(item, dict):
                                parts.extend(str(x) for x in item.values() if isinstance(x, str))
    return "\n".join(p for p in parts if p)


def keyword_hits(text: str, keywords: list[str]) -> list[str]:
    low = (text or "").lower()
    return sorted({k for k in keywords if k.lower() in low})


def classify_live(facility_text: str, page_text: str) -> dict:
    robot = keyword_hits(facility_text, C.ROBOT_KEYWORDS)
    strong = keyword_hits(facility_text, C.AUTOMATION_STRONG_KEYWORDS)
    weak = keyword_hits(facility_text, C.AUTOMATION_WEAK_KEYWORDS)
    staffed = keyword_hits(facility_text, C.STAFFED_KEYWORDS)
    if robot:
        live = C.TYPE_PA
    elif strong:
        live = C.TYPE_AO
    elif facility_text.strip():
        live = C.TYPE_CONV
    else:
        live = "UNKNOWN"
    return {
        "live_type": live,
        "live_robot_kw": "; ".join(robot),
        "live_automation_kw": "; ".join(strong),
        "live_automation_weak_kw": "; ".join(weak),
        "live_staffed_kw": "; ".join(staffed),
        "page_robot_kw_anywhere": "; ".join(keyword_hits(page_text, C.ROBOT_KEYWORDS)),
        "facility_text_len": len(facility_text or ""),
    }

# ============================================================== image URL helpers

_SIZE_SUFFIX_RE = re.compile(
    r"(_(?:R|W|C|Z|M)_\d+_\d+)(?:_(?:R\d+|Q\d+|D|C|M|S|W\d*|H\d*))*(?=\.(?:jpe?g|png|webp|gif)$)", re.I)


def upgrade_image_url(url: str) -> str:
    """Strip Trip.com CDN resize suffixes (e.g. _R_150_150_R5_Q70_D.jpg -> .jpg) to request the
    original. The download step falls back to the URL exactly as Trip.com served it."""
    if not url:
        return url
    base, q = (url.split("?", 1) + [""])[:2]
    up = _SIZE_SUFFIX_RE.sub("", base)
    if q and re.search(r"(resize|proc|imageView|x-oss-process)", q, re.I):
        q = ""
    return up + (("?" + q) if q else "")

# ============================================================== browser session

class Blocked(Exception):
    """Raised when Trip.com serves a CAPTCHA / login wall / rate-limit response."""


class PageCapture:
    """Wraps a Playwright page; records XHR/fetch responses whose URL matches given patterns."""

    def __init__(self, page, patterns: list[str]):
        self.page = page
        self.patterns = patterns
        self._pending = []
        self.parsed: list[tuple[str, Any]] = []
        self.all_xhr: list[tuple[int, str]] = []
        self.block_statuses: list[tuple[int, str]] = []
        page.on("response", self._on_response)

    def _on_response(self, resp) -> None:
        try:
            rtype = resp.request.resource_type
        except Exception:
            rtype = ""
        if rtype not in ("xhr", "fetch"):
            return
        url = resp.url
        self.all_xhr.append((resp.status, url))
        if resp.status in C.BLOCK_HTTP_STATUSES and "trip.com" in url:
            self.block_statuses.append((resp.status, url))
        if any(p in url for p in self.patterns):
            self._pending.append(resp)

    def drain(self) -> list[tuple[str, Any]]:
        """Read bodies of matched responses received since the last drain."""
        new = []
        pending, self._pending = self._pending, []
        for resp in pending:
            try:
                data = resp.json()
            except Exception:
                continue
            new.append((resp.url, data))
        self.parsed.extend(new)
        return new

    def embedded_json(self) -> list:
        js = """() => {
          const out = [];
          const nd = document.getElementById('__NEXT_DATA__');
          if (nd) out.push(nd.textContent);
          for (const s of document.querySelectorAll('script[type="application/json"]')) out.push(s.textContent);
          for (const k of ['__INITIAL_STATE__', '__PRELOADED_STATE__', 'IBU_HOTEL']) {
            try { if (window[k]) out.push(JSON.stringify(window[k])); } catch (e) {}
          }
          return out;
        }"""
        blobs = []
        try:
            for txt in self.page.evaluate(js):
                try:
                    blobs.append(json.loads(txt))
                except Exception:
                    pass
        except Exception:
            pass
        return blobs

    def body_text(self) -> str:
        try:
            return self.page.inner_text("body", timeout=10_000)
        except Exception:
            return ""

    def dom_facility_text(self) -> str:
        js = """() => {
          const els = document.querySelectorAll('[class*="facilit" i],[id*="facilit" i],[class*="amenit" i],[id*="amenit" i],[class*="policy" i]');
          return Array.from(els).map(e => e.innerText).join('\\n');
        }"""
        try:
            return self.page.evaluate(js) or ""
        except Exception:
            return ""

    def check_blocked(self, check_login_text: bool = True) -> str | None:
        if self.block_statuses:
            return f"HTTP_{self.block_statuses[-1][0]}"
        url = (self.page.url or "").lower()
        if any(s in url for s in ("/account/signin", "/login", "passport")):
            return "LOGIN_REQUIRED"
        for fr in self.page.frames:
            if re.search(r"captcha|verify|slider", fr.url or "", re.I):
                return "CAPTCHA"
        text = self.body_text().lower()
        if any(s in text for s in C.BLOCK_TEXT_SIGNALS):
            return "CAPTCHA"
        if check_login_text and any(s in text for s in C.LOGIN_TEXT_SIGNALS):
            return "LOGIN_REQUIRED"
        return None

    def snapshot(self, tag: str) -> None:
        if not C.SAVE_DEBUG_SNAPSHOTS:
            return
        C.DEBUG_DIR.mkdir(parents=True, exist_ok=True)
        try:
            (C.DEBUG_DIR / f"{tag}.html").write_text(self.page.content(), encoding="utf-8")
            (C.DEBUG_DIR / f"{tag}_xhr.txt").write_text(
                "\n".join(f"{s}\t{u}" for s, u in self.all_xhr), encoding="utf-8")
            self.page.screenshot(path=str(C.DEBUG_DIR / f"{tag}.png"), full_page=False)
        except Exception:
            pass


class TripBrowser:
    """One browser + context reused for the whole run (sequential, polite)."""

    def __init__(self, headless: bool | None = None):
        self.headless = C.HEADLESS if headless is None else headless

    def __enter__(self):
        try:
            from playwright.sync_api import sync_playwright
        except ImportError as e:
            raise SystemExit("Playwright is not installed. Run: pip install -r requirements.txt && "
                             "python -m playwright install chromium") from e
        self._pw = sync_playwright().start()
        self.browser = self._pw.chromium.launch(headless=self.headless)
        self.context = self.browser.new_context(
            locale=C.LOCALE, timezone_id=C.TIMEZONE, user_agent=C.USER_AGENT,
            viewport={"width": 1366, "height": 900})
        self.context.set_default_navigation_timeout(C.NAV_TIMEOUT_MS)
        self.context.set_default_timeout(20_000)
        return self

    def __exit__(self, *exc):
        try:
            self.context.close()
            self.browser.close()
        finally:
            self._pw.stop()

    def open(self, url: str, patterns: list[str]) -> PageCapture:
        page = self.context.new_page()
        cap = PageCapture(page, patterns)
        page.goto(url, wait_until="domcontentloaded")
        try:
            page.wait_for_load_state("networkidle", timeout=20_000)
        except Exception:
            pass
        polite_sleep((1.5, 3.0))
        return cap


def save_raw(tag: str, payloads: list[tuple[str, Any]]) -> None:
    if not payloads:
        return
    path = C.RAW_RESPONSE_DIR / f"{tag}.jsonl"
    with path.open("a", encoding="utf-8") as f:
        for url, data in payloads:
            f.write(json.dumps({"ts": now_iso(), "url": url, "data": data}, ensure_ascii=False) + "\n")




# ============================================================== hotel meta from the public review page
#
# 2026-10 live finding: for an automated anonymous browser the hotel *detail* page
# (/hotels/<city>-hotel-detail-<id>/<slug>/) and /hotels/list redirect to /account/signin. That wall is
# not bypassed. The public *review* page (…/review.html) is served normally and its embedded page data
# carries the hotel-level fields this study needs: star, score, review total, city id, brand, a
# displayed price, the top facility tags (names) and the full facility id list (ctripFacilityIds).

FACILITY_NAMES = C.WORK_DIR / "facility_id_names.json"


def _find_review_page_dicts(blobs: list, code: str | None) -> tuple[dict, dict, dict]:
    """(comment-list props, shareData.hotel, shareData.city) for this hotel, {} when absent."""
    cl, sh, city = {}, {}, {}
    for b in blobs:
        for d in iter_dicts(b):
            if code and str(d.get("hotelId")) == str(code):
                if not cl and ("commentRating" in d or "hotelHotTags" in d):
                    cl = d
                if not sh and "ctripFacilityIds" in d:
                    sh = d
            if not city and d.get("type") == "City" and "cityId" in d and "url" in d:
                city = d
    return cl, sh, city


def _remember_facility_names(pairs: list[tuple[int, str]]) -> None:
    if not pairs:
        return
    known = read_json(FACILITY_NAMES, {}) or {}
    changed = False
    for fid, name in pairs:
        if fid is not None and name and str(fid) not in known:
            known[str(fid)] = name
            changed = True
    if changed:
        write_json_atomic(FACILITY_NAMES, known)


def parse_review_page_meta(blobs: list, code: str | None) -> dict | None:
    cl, sh, city = _find_review_page_dicts(blobs, code)
    if not cl and not sh:
        return None
    star = to_number(cl.get("star")) or to_number(sh.get("starId")) or to_number(cl.get("diamond"))
    rating = to_number(cl.get("rating")) or to_number((cl.get("commentRating") or {}).get("ratingAll"))
    full = to_number(cl.get("fullRating")) or to_number((cl.get("commentRating") or {}).get("ratingFull"))
    total = cl.get("totalCount")
    if total is None:
        total = (cl.get("commentRating") or {}).get("showCommentNum")
    dp = cl.get("displayPrice") if isinstance(cl.get("displayPrice"), dict) else {}
    hot_pairs = [(x.get("facilityId"), clean_str(x.get("facilityName")))
                 for x in (cl.get("hotelHotFacilityTags") or []) if isinstance(x, dict)]
    _remember_facility_names(hot_pairs)
    fac_ids = [int(x) for x in (sh.get("ctripFacilityIds") or []) if isinstance(x, (int, float))]
    hot_names = [clean_str(x) for x in (cl.get("hotelHotTags") or []) if isinstance(x, str)]
    basic_names = [clean_str(x.get("name")) for x in (cl.get("hotelFacilityList") or []) if isinstance(x, dict)]
    review_tags = [clean_str(x.get("name")) for x in (cl.get("filterTagList") or [])
                   if isinstance(x, dict) and x.get("filterId") == 0]
    first_reviews = [c.get("content", "") for g in (cl.get("groupList") or []) if isinstance(g, dict)
                     for c in (g.get("commentList") or []) if isinstance(c, dict)]
    brand = cl.get("brand") if isinstance(cl.get("brand"), dict) else {}
    return {
        "trip_hotel_name": clean_str(cl.get("hotelName")) or clean_str(sh.get("enName")),
        "trip_hotel_name_cn": clean_str(sh.get("cnName")),
        "star_rating": star if star and star > 0 else None,
        "star_source": "star" if to_number(cl.get("star")) else ("starId" if to_number(sh.get("starId")) else
                                                                 ("diamond" if to_number(cl.get("diamond")) else "")),
        "overall_rating": rating,
        "overall_rating_scale": full or (10 if rating and rating > 5 else (5 if rating else None)),
        "total_review_count": int(total) if isinstance(total, (int, float)) else None,
        "price_value": to_number(dp.get("price")),
        "price_currency": "USD" if to_number(dp.get("price")) else None,
        "price_checkin": clean_str(dp.get("checkIn")),
        "trip_city_id": int(cl.get("cityId") or sh.get("cityId") or 0) or None,
        "trip_city_name": clean_str(cl.get("cityName")) or clean_str(city.get("name")),
        "trip_city_url": clean_str(city.get("url")),
        "trip_country_name": clean_str(cl.get("countryName")),
        "trip_brand": clean_str(brand.get("name")),
        "open_year": clean_str(cl.get("openYear") or sh.get("openYear")),
        "facility_ids": fac_ids,
        "facility_hot_tags": hot_names,
        "facility_basic": basic_names,
        "review_tags": review_tags,
        "first_page_robot_mentions": sum(1 for t in first_reviews if keyword_hits(t, C.ROBOT_KEYWORDS)),
        "meta_source": "review_page_embedded",
    }


def classify_live_from_meta(meta: dict) -> dict:
    """Live operational type from Trip.com facility data on the public review page.
    Robot evidence = facility id in C.ROBOT_FACILITY_IDS ("Service robots") or a robot keyword in a
    facility tag name. Front-desk automation = id in C.AUTOMATION_FACILITY_IDS or a strong keyword
    in a facility tag name. Review text is recorded as a side signal but never decides the type."""
    ids = set(meta.get("facility_ids") or [])
    names = list(meta.get("facility_hot_tags") or []) + list(meta.get("facility_basic") or [])
    known = read_json(FACILITY_NAMES, {}) or {}
    names_all = names + [known[str(i)] for i in ids if str(i) in known]
    text = "\n".join(names_all)
    robot = keyword_hits(text, C.ROBOT_KEYWORDS)
    if ids & set(C.ROBOT_FACILITY_IDS) and not robot:
        robot = ["facility_id:" + ",".join(str(i) for i in sorted(ids & set(C.ROBOT_FACILITY_IDS)))]
    strong = keyword_hits(text, C.AUTOMATION_STRONG_KEYWORDS)
    auto_ids = ids & set(C.AUTOMATION_FACILITY_IDS)
    if auto_ids and not strong:
        strong = ["facility_id:" + ",".join(str(i) for i in sorted(auto_ids))]
    weak = keyword_hits(text, C.AUTOMATION_WEAK_KEYWORDS)
    staffed = keyword_hits(text, C.STAFFED_KEYWORDS)
    if robot:
        live = C.TYPE_PA
    elif strong:
        live = C.TYPE_AO
    elif ids or names:
        live = C.TYPE_CONV
    else:
        live = "UNKNOWN"
    return {
        "live_type": live,
        "live_robot_kw": "; ".join(robot),
        "live_automation_kw": "; ".join(strong),
        "live_automation_weak_kw": "; ".join(weak),
        "live_staffed_kw": "; ".join(staffed),
        "page_robot_kw_anywhere": "; ".join(keyword_hits("\n".join(meta.get("review_tags") or []), C.ROBOT_KEYWORDS)),
        "facility_text_len": len(text),
    }


def meta_from_capture(cap: "PageCapture", code: str | None, detail_url: str = "") -> dict | None:
    """Parse + cache hotel meta from an already opened review page (used by the review crawler so a
    hotel is opened only once). Returns None when the page carries no hotel data."""
    meta = parse_review_page_meta(cap.embedded_json(), code)
    if not meta:
        return None
    out = {"trip_code": code, "detail_url": detail_url, "fetched_at": now_iso(), **meta}
    out.update(classify_live_from_meta(meta))
    out["fetch_status"] = "OK"
    if code:
        write_json_atomic(C.META_CACHE_DIR / f"{code}.json", out)
    return out


def fetch_hotel_meta(browser: "TripBrowser", detail_url: str, run_cfg: dict, tag: str,
                     force: bool = False) -> dict:
    """Hotel metadata + live facility classification, cached per Trip.com code. Opens the hotel's
    public review page once (light visit: no sorting, no paging). Never retries through a block."""
    code = trip_code_from_url(detail_url)
    cache = C.META_CACHE_DIR / f"{code}.json"
    if not force and cache.exists():
        cached = read_json(cache, {})
        if cached.get("fetch_status") == "OK":
            cached.update(classify_live_from_meta(cached))     # re-apply current keyword / id rules
            return cached
    out = {"trip_code": code, "detail_url": detail_url, "fetched_at": now_iso()}
    cap = None
    try:
        cap = browser.open(review_url_from_detail(detail_url), patterns=["__no_capture__"])
        blocked = cap.check_blocked(check_login_text=False)
        if blocked:
            cap.snapshot(f"meta_blocked_{tag}_{code}")
            out.update(fetch_status="BLOCKED", fetch_error=blocked)
        else:
            meta = meta_from_capture(cap, code, detail_url)
            if meta:
                out = meta
            else:
                cap.snapshot(f"meta_nodata_{tag}_{code}")
                out.update(fetch_status="ERROR", fetch_error="NO_HOTEL_DATA_ON_REVIEW_PAGE")
    except Exception as e:  # noqa: BLE001 - we log and continue
        out.update(fetch_status="ERROR", fetch_error=f"{type(e).__name__}: {e}"[:300])
        if cap:
            cap.snapshot(f"meta_error_{tag}_{code}")
    finally:
        if cap:
            try:
                cap.page.close()
            except Exception:
                pass
    if out.get("fetch_status") != "OK":
        write_json_atomic(cache, out)
    return out
