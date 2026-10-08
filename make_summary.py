"""
Stage 5 — Crawling log (05), sample summary (06) and the terminal summary.

06_sample_summary.xlsx keeps a hotel-level table (one row per hotel with review/photo
aggregates) and computes every summary figure with Excel formulas over that table (COUNTIFS /
SUMIFS / AVERAGEIFS), so numbers are auditable and recalculate if the table is edited.
"""
from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import common as U
import config as C
import crawl_tripcom_reviews as CR
import download_review_images as DL

log = U.get_logger("summary")
GROUPS = [C.GROUP_UNMANNED, C.GROUP_CONVENTIONAL]


def hotel_brands() -> dict[str, str]:
    out = {}
    for path, sheet in ((C.FINAL_UNMANNED, "unmanned_automated_100"),
                        (C.FINAL_CONVENTIONAL if C.FINAL_CONVENTIONAL.exists() else C.DRAFT_CONVENTIONAL, "conventional")):
        try:
            df = pd.read_excel(path, sheet_name=sheet)
            out.update(dict(zip(df.hotel_id, df.brand)))
        except Exception:
            pass
    return out


def build_hotel_level(pairs: pd.DataFrame, reviews: pd.DataFrame, state: dict, man: pd.DataFrame) -> pd.DataFrame:
    brands = hotel_brands()
    rows = []
    for _, h in pairs.iterrows():
        rv = reviews[reviews.hotel_id == h.hotel_id] if len(reviews) else reviews
        st = state.get(h.hotel_id, {})
        m = man[man.hotel_id == h.hotel_id] if len(man) else man
        rated = pd.to_numeric(rv.get("review_rating", pd.Series(dtype=float)), errors="coerce").dropna()
        rows.append({
            "pair_id": h.pair_id, "hotel_id": h.hotel_id, "group": h.group, "operational_type": h.operational_type,
            "hotel_name": h.hotel_name, "country": h.country, "city": h.city,
            "brand": brands.get(h.hotel_id, U.brand_of(h.hotel_name)[0]),
            "star_label": (str(int(h.star_rating)) if pd.notna(h.star_rating) else "unknown"),
            "overall_rating": h.overall_rating,
            "total_review_count_tripcom": st.get("total_review_count_tripcom", h.get("total_review_count")),
            "crawl_status": st.get("status", "NOT_CRAWLED"),
            "n_reviews": len(rv), "n_rated": len(rated), "sum_rating": float(rated.sum()) if len(rated) else 0.0,
            "n_with_images": int(rv["has_image"].astype(str).str.lower().eq("true").sum()) if len(rv) else 0,
            "n_images_attached_upto10": int(pd.to_numeric(rv.get("image_count", 0), errors="coerce").clip(upper=10).sum()) if len(rv) else 0,
            "n_images_downloaded": int((m.status == "OK").sum()) if len(m) else 0,
            "n_images_failed": int((m.status == "FAILED").sum()) if len(m) else 0,
        })
    return pd.DataFrame(rows)


def write_crawl_log(pairs: pd.DataFrame, hl: pd.DataFrame, state: dict) -> pd.DataFrame:
    rows = []
    for _, h in pairs.iterrows():
        st = state.get(h.hotel_id, {})
        r = hl[hl.hotel_id == h.hotel_id].iloc[0]
        rows.append({
            "pair_id": h.pair_id, "hotel_id": h.hotel_id, "group": h.group, "operational_type": h.operational_type,
            "hotel_name": h.hotel_name, "trip_review_url": h.trip_review_url,
            "target_review_count": st.get("target_review_count"),
            "collected_review_count": int(r.n_reviews),
            "reviews_with_images": int(r.n_with_images),
            "total_review_images": int(r.n_images_attached_upto10),
            "downloaded_images": int(r.n_images_downloaded),
            "failed_images": int(r.n_images_failed),
            "crawl_status": st.get("status", "NOT_CRAWLED"),
            "retry_count": st.get("retry_count", 0) if st else None,
            "failure_reason": st.get("failure_reason", "" if st else "not attempted"),
            "crawl_timestamp": st.get("crawl_timestamp"),
            "total_review_count_tripcom": st.get("total_review_count_tripcom"),
            "collected_all_pages": st.get("collected_all"), "sort_mode": st.get("sort_mode"),
            "pages_clicked": st.get("pages_clicked"), "error_detail": st.get("error_detail", ""),
            "pair_match_status": h.get("match_status"),
        })
    df = pd.DataFrame(rows)
    notes = [("generated_at", U.now_iso()),
             ("total_review_images", "photos attached to the sampled reviews, capped at 10 per review"),
             ("status", "SUCCESS = collected ≥ min(100, Trip.com total); PARTIAL = some but fewer; FAILED = none; "
                        "NOT_CRAWLED = not attempted yet"),
             ("failure_reason codes", "CAPTCHA, LOGIN_REQUIRED, HTTP_403/429/430 (blocked, not bypassed); "
                                      "NO_NEXT_PAGE_CONTROL / PAGINATION_EXHAUSTED (Trip.com showed no more reviews to an "
                                      "anonymous visitor); NO_REVIEW_PAYLOAD; TIMEOUT; ERROR; NO_REVIEWS_ON_TRIPCOM"),
             ("VISIBLE_LIMIT_REACHED", "the public review page stops offering 'Show More' after 5 pages x 15 = 75 newest "
                                       "reviews for an anonymous visitor; the limit was respected, not worked around"),
             ("ALL_LISTED_REVIEWS_COLLECTED", "the newest-sorted list ended before 75 (last page short): every review "
                                              "Trip.com lists was collected although its headline review total is larger"),
             ("target_review_count", "min(100, Trip.com review total); with the 75-review visible limit a hotel with more "
                                     "than 75 reviews is PARTIAL by definition")]
    U.write_excel(C.CRAWL_LOG, {"crawling_log": df}, notes)
    return df


def write_summary_workbook(hl: pd.DataFrame) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    hs = wb.create_sheet("hotel_level")
    cols = list(hl.columns)
    hs.append(cols)
    for rec in hl.itertuples(index=False):
        hs.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in rec])
    n = len(hl) + 1
    col = {c: get_column_letter(i + 1) for i, c in enumerate(cols)}

    def rng(c: str) -> str:
        return f"hotel_level!${col[c]}$2:${col[c]}${max(n, 2)}"

    bold, head_fill, sec_fill = Font(name="Arial", bold=True), PatternFill("solid", start_color="DDE4EE"), \
        PatternFill("solid", start_color="F2F2F2")
    ws.append(["Metric", "ALL", C.GROUP_UNMANNED, C.GROUP_CONVENTIONAL, "Definition"])

    def gcrit(g):          # criteria pair for COUNTIFS/SUMIFS; ALL uses wildcard
        return f'{rng("group")},"{g}"' if g else f'{rng("group")},"*"'

    def section(title):
        ws.append([title])
        ws.cell(ws.max_row, 1).font = bold
        for c in range(1, 6):
            ws.cell(ws.max_row, c).fill = sec_fill

    def metric(label, fn, definition, fmt=None):
        r = ws.max_row + 1
        ws.append([label] + [fn(g, r) for g in (None, *GROUPS)] + [definition])
        if fmt:
            for c in range(2, 5):
                ws.cell(r, c).number_format = fmt
        return r

    section("[호텔 표본]")
    r_hot = metric("총 호텔 수", lambda g, r: f"=COUNTIFS({gcrit(g)})", "hotels in the current pair plan")
    metric("PHYSICAL_AI 호텔 수", lambda g, r: f'=COUNTIFS({gcrit(g)},{rng("operational_type")},"{C.TYPE_PA}")', "")
    metric("AUTOMATION_ONLY 호텔 수", lambda g, r: f'=COUNTIFS({gcrit(g)},{rng("operational_type")},"{C.TYPE_AO}")', "")
    metric("CONVENTIONAL 호텔 수", lambda g, r: f'=COUNTIFS({gcrit(g)},{rng("operational_type")},"{C.TYPE_CONV}")', "")
    metric("평균 Trip.com 총 리뷰 수 / 호텔", lambda g, r: f'=IFERROR(AVERAGEIFS({rng("total_review_count_tripcom")},{gcrit(g)}),"n/a")',
           "Trip.com-reported review total (not only collected)", "0.0")
    metric("평균 호텔 overall rating", lambda g, r: f'=IFERROR(AVERAGEIFS({rng("overall_rating")},{gcrit(g)}),"n/a")',
           "Trip.com hotel score (scale as displayed, usually /10)", "0.00")
    section("[리뷰]")
    r_rev = metric("전체 수집 리뷰 수", lambda g, r: f'=SUMIFS({rng("n_reviews")},{gcrit(g)})', "rows in 04 CSV")
    metric("평균 review rating", lambda g, r: f'=IFERROR(SUMIFS({rng("sum_rating")},{gcrit(g)})/SUMIFS({rng("n_rated")},{gcrit(g)}),"n/a")',
           "mean of review_rating (Trip.com review scale 0–10)", "0.00")
    metric("평균 수집 리뷰 수 / 호텔", lambda g, r: f'=IFERROR(AVERAGEIFS({rng("n_reviews")},{gcrit(g)}),"n/a")', "", "0.0")
    for st in ("SUCCESS", "PARTIAL", "FAILED", "NOT_CRAWLED"):
        metric(f"크롤링 {st} 호텔 수", lambda g, r, st=st: f'=COUNTIFS({gcrit(g)},{rng("crawl_status")},"{st}")', "")
    section("[사진]")
    r_wi = metric("사진이 있는 리뷰 수", lambda g, r: f'=SUMIFS({rng("n_with_images")},{gcrit(g)})', "")
    metric("사진이 있는 리뷰 비율", lambda g, r: f'=IFERROR({get_column_letter(2 + (GROUPS.index(g) + 1 if g else 0))}{r_wi}/'
                                            f'{get_column_letter(2 + (GROUPS.index(g) + 1 if g else 0))}{r_rev},"n/a")', "", "0.0%")
    r_dl = metric("다운로드된 사진 수", lambda g, r: f'=SUMIFS({rng("n_images_downloaded")},{gcrit(g)})', "")
    metric("다운로드 실패 사진 수", lambda g, r: f'=SUMIFS({rng("n_images_failed")},{gcrit(g)})', "")
    metric("리뷰당 평균 사진 수 (전체 리뷰 기준)", lambda g, r: f'=IFERROR({get_column_letter(2 + (GROUPS.index(g) + 1 if g else 0))}{r_dl}/'
                                                    f'{get_column_letter(2 + (GROUPS.index(g) + 1 if g else 0))}{r_rev},"n/a")', "", "0.00")
    metric("사진 리뷰당 평균 사진 수", lambda g, r: f'=IFERROR({get_column_letter(2 + (GROUPS.index(g) + 1 if g else 0))}{r_dl}/'
                                          f'{get_column_letter(2 + (GROUPS.index(g) + 1 if g else 0))}{r_wi},"n/a")', "", "0.00")
    _ = r_hot

    def breakdown(title, keys: list[tuple], key_cols: list[str], share=False):
        s = wb.create_sheet(title)
        s.append(key_cols + ["ALL", C.GROUP_UNMANNED, C.GROUP_CONVENTIONAL] + (["share_of_unmanned"] if share else []))
        for i, key in enumerate(keys, start=2):
            crit = ",".join(f'{rng(kc)},"{kv}"' for kc, kv in zip(key_cols, key))
            row = list(key) + [f"=COUNTIFS({crit})"] + [f'=COUNTIFS({crit},{rng("group")},"{g}")' for g in GROUPS]
            k = len(key_cols)
            if share:
                tot = f"SUM(${get_column_letter(k + 2)}$2:${get_column_letter(k + 2)}${len(keys) + 1})"
                row.append(f"=IFERROR({get_column_letter(k + 2)}{i}/{tot},0)")
            s.append(row)
            if share:
                s.cell(i, len(row)).number_format = "0.0%"
        return s

    breakdown("by_country", sorted({(c,) for c in hl.country}), ["country"])
    breakdown("by_city", sorted(set(zip(hl.country, hl.city))), ["country", "city"])
    brand_order = hl[hl.group == C.GROUP_UNMANNED].brand.value_counts().index.tolist()
    brand_order += sorted(set(hl.brand) - set(brand_order))
    breakdown("by_brand", [(b,) for b in brand_order], ["brand"], share=True)
    breakdown("by_star", sorted({(s,) for s in hl.star_label}), ["star_label"])

    notes = wb.create_sheet("Notes")
    for row in [("item", "detail"), ("generated_at", U.now_iso()),
                ("data", "hotel_level sheet is generated from work/pairs_current.csv, work/crawl_state.json, "
                         "04_tripcom_reviews_raw.csv and work/image_manifest.csv; all Summary/by_* figures are formulas"),
                ("brand concentration", "by_brand.share_of_unmanned = share of the 100 unmanned/automated hotels"),
                ("rating scales", "review_rating is Trip.com's 0–10 review scale; hotel overall rating is as displayed"),
                ("ALL column", "uses group criterion '*' (all hotels)")]:
        notes.append(list(row))

    for sheet in wb.worksheets:
        for r in sheet.iter_rows():
            for c in r:
                c.font = Font(name="Arial", size=10, bold=(c.row == 1) or (c.font and c.font.bold))
                if c.row == 1:
                    c.fill = head_fill
                    c.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.freeze_panes = "A2" if sheet.title != "Summary" else "B2"
        for i, column in enumerate(sheet.columns, 1):
            width = max(len(str(c.value)) if c.value is not None and not str(c.value).startswith("=") else 10
                        for c in column)
            sheet.column_dimensions[get_column_letter(i)].width = max(10, min(55, width + 2))
    wb.calculation.fullCalcOnLoad = True
    wb.save(C.SAMPLE_SUMMARY)


def terminal_summary(hl: pd.DataFrame) -> str:
    un = hl[hl.group == C.GROUP_UNMANNED]
    cv = hl[hl.group == C.GROUP_CONVENTIONAL]
    st = hl.crawl_status.value_counts()
    pairs_done = int((hl.groupby("pair_id").group.nunique() == 2).sum())
    lines = [
        "-----------------------------------",
        f"UNMANNED / AUTOMATED HOTELS: {len(un)}",
        f"PHYSICAL AI: {(un.operational_type == C.TYPE_PA).sum()}",
        f"AUTOMATION ONLY: {(un.operational_type == C.TYPE_AO).sum()}",
        f"CONVENTIONAL HOTELS: {len(cv)}",
        f"TOTAL HOTELS: {len(hl)}",
        f"COMPLETE PAIRS: {pairs_done}",
        "",
        f"SUCCESSFUL HOTEL CRAWLS: {st.get('SUCCESS', 0)}",
        f"PARTIAL: {st.get('PARTIAL', 0)}",
        f"FAILED: {st.get('FAILED', 0)}" + (f"   (NOT CRAWLED: {st.get('NOT_CRAWLED', 0)})" if st.get("NOT_CRAWLED") else ""),
        "",
        f"UNMANNED / AUTOMATED REVIEWS: {int(un.n_reviews.sum())}",
        f"CONVENTIONAL REVIEWS: {int(cv.n_reviews.sum())}",
        f"TOTAL REVIEWS: {int(hl.n_reviews.sum())}",
        "",
        f"REVIEWS WITH PHOTOS: {int(hl.n_with_images.sum())}",
        f"TOTAL PHOTOS DOWNLOADED: {int(hl.n_images_downloaded.sum())}",
        f"PHOTO DOWNLOAD FAILURES: {int(hl.n_images_failed.sum())}",
        "-----------------------------------",
    ]
    return "\n".join(lines)


def run() -> str:
    U.ensure_dirs()
    pairs = pd.read_csv(C.PAIRS_CURRENT_CSV, dtype={"hotel_id": str})
    reviews = (pd.read_csv(C.REVIEWS_CSV, dtype={"review_id": str, "hotel_id": str}, encoding="utf-8-sig")
               if C.REVIEWS_CSV.exists() else pd.DataFrame(columns=["hotel_id", "review_rating", "has_image", "image_count"]))
    state = CR.load_state()
    man = DL.load_manifest()
    if len(man):
        man = man.drop_duplicates(["review_id", "photo_index"], keep="last")
    hl = build_hotel_level(pairs, reviews, state, man)
    write_crawl_log(pairs, hl, state)
    write_summary_workbook(hl)
    text = terminal_summary(hl)
    print(text)
    log.info("Summary written → %s, %s", C.CRAWL_LOG.name, C.SAMPLE_SUMMARY.name)
    return text


if __name__ == "__main__":
    run()
